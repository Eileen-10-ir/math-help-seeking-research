"""Readable Chinese Word brief, using the managed document runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / '团队专属交接20261014/lbb_更新任务包/01_给本人阅读'
source = (BASE / '项目回顾与本次任务.md').read_text(encoding='utf-8')
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.8)
for style_name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
    style = doc.styles[style_name]
    style.font.name = 'SimSun'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.2
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Title'].font.size = Pt(20)
doc.styles['Heading 1'].font.size = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(5)
lines = source.splitlines()
index = 0
while index < len(lines):
    line = lines[index].strip()
    if not line:
        index += 1; continue
    if line.startswith('|'):
        table_lines = []
        while index < len(lines) and lines[index].strip().startswith('|'):
            table_lines.append(lines[index].strip());index += 1
        cells = [[s.strip() for s in row.strip('|').split('|')] for row in table_lines if not set(row.replace('|', '').replace(':', '').replace('-', '').strip()) == set()]
        table = doc.add_table(rows=1, cols=len(cells[0]))
        table.style = 'Table Grid'
        for j, value in enumerate(cells[0]):
            table.rows[0].cells[j].text = value
        for row in cells[1:]:
            new = table.add_row()
            for j, value in enumerate(row):
                new.cells[j].text = value
        header = OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(header)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
        continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
        doc.add_paragraph('2026年10月14日  供组长确认后交接', 'Normal')
    elif line.startswith('## '):
        p = doc.add_paragraph(line[3:], 'Heading 1')
        p.paragraph_format.keep_with_next = True
    else:
        p = doc.add_paragraph(line.replace('`', ''), 'Normal')
    index += 1
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('数学单科研究交接  ')
field = OxmlElement('w:fldSimple');field.set(qn('w:instr'), 'PAGE');footer._p.append(field)
for root in (doc.styles.element, doc.element):
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
out = BASE / '项目回顾与本次任务.docx'
doc.save(out)
print(out)
